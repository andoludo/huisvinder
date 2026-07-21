from decimal import Decimal
from pathlib import Path

import openpyxl
import sqlmodel

from huisvinder.database.crud import HuisVinderDb
from huisvinder.database.schemas import PropertySalesRecordORM
from huisvinder.services import add_property_sales_record, load_property_sales_records

SEGMENTS = [
    "Alle huizen met 2, 3, 4 of meer gevels (excl. appartementen)",
    "Huizen met 2 of 3 gevels (gesloten + halfopen bebouwing)",
    "Huizen met 4 of meer gevels (open bebouwing)",
    "Appartementen",
]
METRICS = ["aantal transacties", "mediaan prijs(€)", "eerste kwartiel prijs(€)", "derde kwartiel prijs(€)"]

# metric blocks start at F, K, P, U with a blank spacer column before each
SEGMENT_START_COLUMNS = [6, 11, 16, 21]


def write_workbook(path: Path, rows: list[tuple[tuple, list]]) -> Path:
    """Build a fixture workbook mirroring the Statbel layout.

    Each row is ((refnis, lokaliteit, jaar, periode), [16 metric values in
    segment order]). Includes the banner rows, spacer columns and the trailing
    helper-formula column of the real export."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Per gemeente"
    sheet.cell(row=1, column=1, value="Vastgoedcategorieën op basis van aard grond")
    for segment_name, start in zip(SEGMENTS, SEGMENT_START_COLUMNS, strict=True):
        sheet.cell(row=1, column=start, value="verkopen van onroerende goederen")
        sheet.cell(row=2, column=start, value=segment_name)
        for offset, metric in enumerate(METRICS):
            sheet.cell(row=3, column=start + offset, value=metric)
    for column, name in enumerate(["refnis", "lokaliteit", "jaar", "periode"], start=1):
        sheet.cell(row=3, column=column, value=name)
    sheet.cell(row=3, column=28, value=0)
    for row_number, (keys, metrics) in enumerate(rows, start=4):
        for column, value in enumerate(keys, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        for segment_index, start in enumerate(SEGMENT_START_COLUMNS):
            for offset in range(4):
                value = metrics[segment_index * 4 + offset]
                if value is not None:
                    sheet.cell(row=row_number, column=start + offset, value=value)
        sheet.cell(row=row_number, column=28, value="=IF(A4=A3,AB3,AB3+1)")
    workbook.save(path)
    return path


FULL_ROW = (("11001", "AARTSELAAR", "2010", "Q1"), [29, 252000, 225000, 290000] * 4)
SUPPRESSED_ROW = (("23038", "HOEILAART", "2011", "Q3"), [4, None, None, None] + [29, 252000, 225000, 290000] * 3)


def load_orm_rows(db: HuisVinderDb) -> list[PropertySalesRecordORM]:
    with sqlmodel.Session(db._engine) as session:
        return list(session.exec(sqlmodel.select(PropertySalesRecordORM)).all())


def test_thousands_separator_is_stripped(tmp_path: Path):
    # comma is a thousands separator in this export: "300,000" is 300000, not 300
    workbook = write_workbook(
        tmp_path / "sales.xlsx",
        [(("11001", "AARTSELAAR", "2010", "Q1"), ["1,234", "300,000", "225,000", "310,000"] + [None] * 12)],
    )
    (record,) = load_property_sales_records(workbook)
    assert record.all_houses_transactions == 1234
    assert record.all_houses_median == Decimal(300000)
    assert record.all_houses_q1 == Decimal(225000)
    assert record.all_houses_q3 == Decimal(310000)


def test_suppressed_quartiles_load_as_null(tmp_path: Path):
    workbook = write_workbook(tmp_path / "sales.xlsx", [SUPPRESSED_ROW])
    db_path = tmp_path / "test.db"
    add_property_sales_record(db_path, workbook)
    (row,) = load_orm_rows(HuisVinderDb(database_path=db_path))
    assert row.nis_code == "23038"
    assert row.all_houses_transactions == 4
    assert row.all_houses_median is None
    assert row.all_houses_q1 is None
    assert row.all_houses_q3 is None
    assert row.terraced_semi_median == 252000


def test_reload_is_idempotent(tmp_path: Path):
    workbook = write_workbook(tmp_path / "sales.xlsx", [FULL_ROW, SUPPRESSED_ROW])
    db_path = tmp_path / "test.db"
    add_property_sales_record(db_path, workbook)
    add_property_sales_record(db_path, workbook)
    assert len(load_orm_rows(HuisVinderDb(database_path=db_path))) == 2


def test_invalid_row_is_skipped_and_logged(tmp_path: Path, caplog):
    bad_row = (("11002", "ANTWERPEN", "2010", "Q9"), [29, 252000, 225000, 290000] * 4)
    workbook = write_workbook(tmp_path / "sales.xlsx", [FULL_ROW, bad_row])
    with caplog.at_level("WARNING"):
        records = load_property_sales_records(workbook)
    assert [record.nis_code for record in records] == ["11001"]
    assert "refnis=11002" in caplog.text
    assert "Q9" in caplog.text
